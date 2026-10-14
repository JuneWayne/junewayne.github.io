"""Export a rendered slide PDF plus PPTX animation placements for the static viewer.
Usage: python export-parc-slides.py deck.pptx rendered.pdf
Requires python-pptx, PyMuPDF and Pillow. Render the same deck with LibreOffice first.
"""
import json
import sys
from pathlib import Path
from zipfile import ZipFile
import pymupdf as fitz
from PIL import Image
from pptx import Presentation

pptx, pdf = map(Path, sys.argv[1:])
root = Path(__file__).resolve().parents[1]
out = root / 'assets/slides/parc'
out.mkdir(parents=True, exist_ok=True)
presentation = Presentation(pptx)
document = fitz.open(pdf)
assert len(document) == len(presentation.slides) == 83
slides = []
media = set()
with ZipFile(pptx) as archive:
    for index, slide in enumerate(presentation.slides):
        number = index + 1
        paragraphs = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        for shape in slide.shapes:
            if shape.has_table:
                paragraphs.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape and title_shape.text.strip() else (paragraphs[0].split('\n')[0] if paragraphs else f'Slide {number}')
        if number == 1:
            title = 'Shear Flow Modeling'
        page = document[index]
        bitmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        background = f'slide-{number:02}.webp'
        Image.frombytes('RGB', [bitmap.width, bitmap.height], bitmap.samples).save(out / background, quality=90, method=6)
        animations = []
        foregrounds = []
        seen_animation = False
        for shape in slide.shapes:
            if shape.shape_type != 13:
                continue
            animated = shape.image.ext == 'gif'
            if animated:
                seen_animation = True
            elif not seen_animation:
                continue
            part = slide.part.related_part(shape._pic.blipFill.blip.rEmbed)
            filename = Path(str(part.partname)).name
            data = archive.read(str(part.partname).lstrip('/'))
            if filename not in media:
                (out / filename).write_bytes(data)
                media.add(filename)
            (animations if animated else foregrounds).append({
                'src': filename,
                'x': shape.left / presentation.slide_width * 100,
                'y': shape.top / presentation.slide_height * 100,
                'width': shape.width / presentation.slide_width * 100,
                'height': shape.height / presentation.slide_height * 100,
                'crop': [shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom],
                'rotation': shape.rotation,
            })
        slides.append({'title': title, 'background': background, 'text': '\n\n'.join(paragraphs), 'animations': animations, 'foregrounds': foregrounds})
(out / 'slides.json').write_text(json.dumps({'aspectRatio': presentation.slide_width / presentation.slide_height, 'slides': slides}, indent=2))
print(f'Exported {len(slides)} slides and {len(media)} unique overlay images.')
